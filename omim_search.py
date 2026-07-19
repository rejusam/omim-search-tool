"""Search the OMIM API by keyword and save the results.

Run this file directly to use it. See README.md for the operator guide.
"""

import configparser
import csv
import datetime
import json
import pathlib
import re
import sys
import time
import requests

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


ENTRY_FIELDS = [
    "rank", "mim_number", "prefix", "preferred_title", "status",
    "gene_symbols", "approved_gene_symbol", "gene_name", "cyto_location",
    "chromosome", "phenotype_count", "phenotypes", "inheritance", "omim_url",
]
PHENOTYPE_FIELDS = [
    "mim_number", "preferred_title", "gene_symbols", "approved_gene_symbol",
    "cyto_location", "phenotype", "phenotype_mim_number",
    "phenotype_mapping_key", "phenotype_inheritance",
    "phenotypic_series_number", "entry_url", "phenotype_url",
]
# Columns whose values are gene symbols and must be stored as text so Excel
# does not convert names like SEPT9 or MARCH1 into dates.
TEXT_COLUMNS = {"gene_symbols", "approved_gene_symbol"}
# Columns whose values are OMIM web links and should be written as clickable
# hyperlinks rather than plain text.
URL_COLUMNS = {"omim_url", "entry_url", "phenotype_url"}


class OmimError(Exception):
    """Something went wrong talking to the OMIM API."""


class AuthFailed(OmimError):
    """The API key was rejected (HTTP 401)."""


class QuotaExhausted(OmimError):
    """The API key's request quota is used up (HTTP 429)."""


class OmimClient:
    """A small, throttled, retrying client for the OMIM API."""

    def __init__(self, api_key, base_url="https://api.omim.org/api",
                 session=None, pause_seconds=0.3, sleep=time.sleep):
        self.api_key = api_key
        self.base_url = base_url.rstrip("/")
        if session is None:
            session = requests.Session()
        self.session = session
        self.pause_seconds = pause_seconds
        self.sleep = sleep

    def request(self, path, params):
        """Make one GET request and return the parsed 'omim' payload.

        Retries server errors and real network errors (connection failures,
        timeouts) up to 3 attempts total. Raises AuthFailed on 401 and
        QuotaExhausted on 429 (no retry on either). A 404 means the data does
        not exist, so it returns an empty payload rather than retrying.
        """
        url = self.base_url + "/" + path
        headers = {
            "ApiKey": self.api_key,
            "Accept-Encoding": "gzip",
            "Accept": "application/json",
        }

        max_attempts = 3
        attempt = 1
        while True:
            self.sleep(self.pause_seconds)
            try:
                response = self.session.get(
                    url, params=params, headers=headers, timeout=60
                )
            except requests.exceptions.RequestException as exc:
                if attempt >= max_attempts:
                    raise OmimError(
                        "A network error occurred talking to OMIM."
                    ) from exc
                self.sleep(2 ** attempt)  # 2s, then 4s
                attempt = attempt + 1
                continue

            status = response.status_code

            if status == 200:
                body = response.json()
                return body.get("omim", {})
            if status == 401:
                raise AuthFailed("The API key was rejected.")
            if status == 429:
                raise QuotaExhausted("The API key's request quota is exhausted.")
            if status == 404:
                return {}

            # 400, 500 and anything else: retry a few times, then give up.
            if attempt >= max_attempts:
                raise OmimError("OMIM returned HTTP " + str(status) + ".")
            self.sleep(2 ** attempt)  # 2s, then 4s
            attempt = attempt + 1

    def search(self, query, start, limit, include=None):
        """Run an entry search and return the 'omim' payload."""
        params = {
            "search": query,
            "start": start,
            "limit": limit,
            "format": "json",
        }
        if include is not None:
            params["include"] = include
        return self.request("entry/search", params)


def load_api_key(config_path):
    """Return the saved API key, or None if there is no usable key."""
    parser = configparser.ConfigParser()
    read_files = parser.read(config_path)
    if not read_files:
        return None
    key = parser.get("omim", "api_key", fallback="")
    key = key.strip()
    if key == "":
        return None
    return key


def save_api_key(config_path, api_key):
    """Write the API key to the config file, creating it if needed."""
    parser = configparser.ConfigParser()
    parser["omim"] = {"api_key": api_key.strip()}
    with open(config_path, "w", encoding="utf-8") as config_file:
        parser.write(config_file)


def quote_term(term):
    """Wrap a term in quotes if it contains a space, so OMIM treats it as a phrase."""
    term = term.strip()
    if " " in term:
        return '"' + term + '"'
    return term


def build_query(any_of, must_include, exclude):
    """Build an OMIM query string from three lists of terms.

    any_of: entries may contain any of these (an OR group, but required as a group)
    must_include: entries must contain each of these
    exclude: entries must not contain any of these
    """
    parts = []

    clean_any_of = []
    for term in any_of:
        term = term.strip()
        if term != "":
            clean_any_of.append(quote_term(term))
    if clean_any_of:
        parts.append("+(" + " ".join(clean_any_of) + ")")

    for term in must_include:
        term = term.strip()
        if term != "":
            parts.append("+" + quote_term(term))

    for term in exclude:
        term = term.strip()
        if term != "":
            parts.append("-" + quote_term(term))

    if not parts:
        raise ValueError("Enter at least one search word.")

    return " ".join(parts)


def page_starts(total_results, max_results, page_size=20):
    """Return the list of 'start' offsets needed to fetch the results.

    total_results: how many results OMIM says there are
    max_results: cap on how many to fetch, or None for all
    """
    wanted = total_results
    if max_results is not None and max_results < wanted:
        wanted = max_results

    starts = []
    start = 0
    while start < wanted:
        starts.append(start)
        start = start + page_size
    return starts


def probe_count(client, query):
    """Ask OMIM how many results a query has, cheaply (no includes)."""
    payload = client.search(query, start=0, limit=1)
    search_response = payload.get("searchResponse", {})
    return search_response.get("totalResults", 0)


def fetch_entries(client, query, max_results):
    """Fetch entries for a query, paging with include=geneMap.

    Returns (entries, raw_responses, complete). If a quota stop or repeated
    error interrupts paging, returns whatever was gathered with complete=False.
    """
    # Make first request with include=geneMap to get both totalResults and first page
    try:
        payload = client.search(query, start=0, limit=20, include="geneMap")
    except (QuotaExhausted, OmimError):
        return [], [], False

    search_response = payload.get("searchResponse", {})
    total = search_response.get("totalResults", 0)

    # Get all page start offsets and determine which ones we still need
    starts = page_starts(total, max_results)

    entries = []
    raw_responses = [payload]
    complete = True

    # Extract entries from first page
    entry_list = search_response.get("entryList", [])
    for item in entry_list:
        entries.append(item.get("entry", {}))

    # Fetch remaining pages (skip start=0 since we already have it)
    for start in starts:
        if start == 0:
            continue
        try:
            payload = client.search(query, start=start, limit=20, include="geneMap")
        except (QuotaExhausted, OmimError):
            complete = False
            break
        raw_responses.append(payload)
        search_response = payload.get("searchResponse", {})
        entry_list = search_response.get("entryList", [])
        for item in entry_list:
            entries.append(item.get("entry", {}))

    return entries, raw_responses, complete


def _field(mapping, key):
    """Return mapping[key] as a safe value: '' for a missing key or a JSON null.

    OMIM sometimes includes a field with a null value rather than omitting it,
    so a plain .get(key, "") can still return None. This keeps None out of the
    output cells and out of the ' | ' joins.
    """
    value = mapping.get(key)
    if value is None:
        return ""
    return value


def gene_map_of(entry):
    """Return the dict that holds this entry's gene and location fields, or {}.

    OMIM returns two shapes. Gene entries (prefix *, %, +) carry a 'geneMap'
    dict. Phenotype entries (prefix #) have no geneMap; instead they keep the
    gene and location fields inside their first phenotype map. This returns
    whichever is present so the Entries sheet can show gene columns for both.
    """
    gene_map = entry.get("geneMap")
    if isinstance(gene_map, dict) and gene_map:
        return gene_map
    phenotype_map_list = entry.get("phenotypeMapList")
    if phenotype_map_list:
        return phenotype_map_list[0].get("phenotypeMap", {})
    return {}


def _phenotype_maps_of(entry):
    """Return the list of phenotype map dicts for an entry (possibly empty).

    Phenotype entries hold the list directly on the entry as
    'phenotypeMapList'. Gene entries hold it nested inside their 'geneMap'.
    This collects from whichever place it lives in.
    """
    phenotype_map_list = entry.get("phenotypeMapList")
    if not phenotype_map_list:
        gene_map = entry.get("geneMap")
        if isinstance(gene_map, dict):
            phenotype_map_list = gene_map.get("phenotypeMapList")
    if not phenotype_map_list:
        return []
    result = []
    for item in phenotype_map_list:
        result.append(item.get("phenotypeMap", {}))
    return result


def _gene_field(entry, phenotype_map, key):
    """Look up one gene or location field for a phenotype row.

    Phenotype entries embed the gene fields inside each phenotype map, so a
    row prefers its own phenotype map's value. Gene entries keep those fields
    on the geneMap while the nested phenotype map only carries phenotype data,
    so fall back to the entry's geneMap when the phenotype map has no value.
    """
    value = _field(phenotype_map, key)
    if value != "":
        return value
    gene_map = entry.get("geneMap")
    if isinstance(gene_map, dict):
        return _field(gene_map, key)
    return ""


def entry_to_row(entry, rank):
    """Flatten one entry into a Sheet 1 (Entries) row."""
    mim_number = _field(entry, "mimNumber")
    titles = entry.get("titles", {})
    gene_map = gene_map_of(entry)
    phenotype_maps = _phenotype_maps_of(entry)

    phenotype_names = []
    inheritance_values = []
    for phenotype_map in phenotype_maps:
        name = _field(phenotype_map, "phenotype")
        if name != "":
            phenotype_names.append(name)
        inheritance = _field(phenotype_map, "phenotypeInheritance")
        if inheritance != "" and inheritance not in inheritance_values:
            inheritance_values.append(inheritance)

    if mim_number == "":
        omim_url = ""
    else:
        omim_url = "https://omim.org/entry/" + str(mim_number)

    row = {
        "rank": rank,
        "mim_number": mim_number,
        "prefix": _field(entry, "prefix"),
        "preferred_title": _field(titles, "preferredTitle"),
        "status": _field(entry, "status"),
        "gene_symbols": _field(gene_map, "geneSymbols"),
        "approved_gene_symbol": _field(gene_map, "approvedGeneSymbols"),
        "gene_name": _field(gene_map, "geneName"),
        "cyto_location": _field(gene_map, "cytoLocation"),
        "chromosome": _field(gene_map, "chromosomeSymbol"),
        "phenotype_count": len(phenotype_names),
        "phenotypes": " | ".join(phenotype_names),
        "inheritance": " | ".join(inheritance_values),
        "omim_url": omim_url,
    }
    return row


def entry_to_phenotype_rows(entry):
    """Flatten one entry into zero or more Sheet 2 (Phenotypes) rows."""
    mim_number = _field(entry, "mimNumber")
    titles = entry.get("titles", {})
    phenotype_maps = _phenotype_maps_of(entry)

    if mim_number == "":
        entry_url = ""
    else:
        entry_url = "https://omim.org/entry/" + str(mim_number)

    rows = []
    for phenotype_map in phenotype_maps:
        phenotype_mim = _field(phenotype_map, "phenotypeMimNumber")
        if phenotype_mim == "":
            phenotype_url = ""
        else:
            phenotype_url = "https://omim.org/entry/" + str(phenotype_mim)
        row = {
            "mim_number": mim_number,
            "preferred_title": _field(titles, "preferredTitle"),
            "gene_symbols": _gene_field(entry, phenotype_map, "geneSymbols"),
            "approved_gene_symbol": _gene_field(entry, phenotype_map, "approvedGeneSymbols"),
            "cyto_location": _gene_field(entry, phenotype_map, "cytoLocation"),
            "phenotype": _field(phenotype_map, "phenotype"),
            "phenotype_mim_number": phenotype_mim,
            "phenotype_mapping_key": _field(phenotype_map, "phenotypeMappingKey"),
            "phenotype_inheritance": _field(phenotype_map, "phenotypeInheritance"),
            "phenotypic_series_number": _field(phenotype_map, "phenotypicSeriesNumber"),
            "entry_url": entry_url,
            "phenotype_url": phenotype_url,
        }
        rows.append(row)
    return rows


def slugify(text):
    """Turn a query into a short, filesystem-safe slug."""
    lowered = text.lower()
    replaced = re.sub(r"[^a-z0-9]+", "_", lowered)
    trimmed = replaced.strip("_")
    if len(trimmed) > 40:
        trimmed = trimmed[:40].strip("_")
    return trimmed


def make_run_folder(results_dir, query):
    """Create and return a unique folder for this run's output."""
    results_dir = pathlib.Path(results_dir)
    timestamp = datetime.datetime.now().strftime("%Y%m%d-%H%M")
    slug = slugify(query)
    base_name = "omim_" + slug + "_" + timestamp

    folder = results_dir / base_name
    suffix = 2
    while folder.exists():
        folder = results_dir / (base_name + "_" + str(suffix))
        suffix = suffix + 1

    folder.mkdir(parents=True)
    return folder


def write_csv(path, rows, fieldnames):
    """Write rows to a UTF-8-with-BOM csv so Excel reads it correctly."""
    with open(path, "w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def write_json(path, raw_responses, metadata):
    """Write the raw API responses plus a metadata block."""
    document = {"metadata": metadata, "responses": raw_responses}
    with open(path, "w", encoding="utf-8") as json_file:
        json.dump(document, json_file, indent=2)


def _write_sheet(sheet, fieldnames, rows):
    """Write a header row and data rows, protecting gene-symbol columns."""
    for column_index, name in enumerate(fieldnames, start=1):
        sheet.cell(row=1, column=column_index, value=name)

    row_index = 2
    for row in rows:
        for column_index, name in enumerate(fieldnames, start=1):
            value = row.get(name, "")
            cell = sheet.cell(row=row_index, column=column_index)
            if name in TEXT_COLUMNS:
                cell.value = str(value)
                cell.number_format = "@"
            elif name in URL_COLUMNS and value != "":
                cell.value = value
                cell.hyperlink = value
                cell.style = "Hyperlink"
            else:
                cell.value = value
        row_index = row_index + 1

    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(fieldnames))
    last_row = max(row_index - 1, 1)
    sheet.auto_filter.ref = "A1:" + last_column + str(last_row)


def write_xlsx(path, entry_rows, phenotype_rows, search_info):
    """Write the three-sheet results workbook."""
    workbook = Workbook()

    entries_sheet = workbook.active
    entries_sheet.title = "Entries"
    _write_sheet(entries_sheet, ENTRY_FIELDS, entry_rows)

    phenotypes_sheet = workbook.create_sheet("Phenotypes")
    _write_sheet(phenotypes_sheet, PHENOTYPE_FIELDS, phenotype_rows)

    info_sheet = workbook.create_sheet("Search info")
    info_sheet.cell(row=1, column=1, value="Field")
    info_sheet.cell(row=1, column=2, value="Value")
    info_row = 2
    for label, value in search_info.items():
        info_sheet.cell(row=info_row, column=1, value=label)
        info_sheet.cell(row=info_row, column=2, value=str(value))
        info_row = info_row + 1

    workbook.save(path)


def write_run(folder, entry_rows, phenotype_rows, raw_responses, search_info):
    """Write all four output files into the run folder."""
    folder = pathlib.Path(folder)
    write_xlsx(folder / "results.xlsx", entry_rows, phenotype_rows, search_info)
    write_csv(folder / "entries.csv", entry_rows, ENTRY_FIELDS)
    write_csv(folder / "phenotypes.csv", phenotype_rows, PHENOTYPE_FIELDS)
    write_json(folder / "raw.json", raw_responses, search_info)


def split_terms(text):
    """Split a comma-separated line into a list of non-empty, trimmed terms."""
    terms = []
    for piece in text.split(","):
        piece = piece.strip()
        if piece != "":
            terms.append(piece)
    return terms


def build_search_info(query, mode, total, entry_rows, phenotype_rows, complete):
    """Assemble the Sheet 3 / metadata dictionary."""
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    if complete:
        complete_text = "yes"
    else:
        complete_text = "no - stopped early (quota or error); results are partial"
    info = {
        "Query sent to OMIM": query,
        "Search mode": mode,
        "Run at": timestamp,
        "Total results reported by OMIM": total,
        "Entries written (Sheet 1)": len(entry_rows),
        "Phenotype rows written (Sheet 2)": len(phenotype_rows),
        "Run complete": complete_text,
        "Licence": (
            "Contains OMIM data (c) Johns Hopkins University. Not for "
            "redistribution. Do not forward these files onward without "
            "checking OMIM's terms."
        ),
    }
    return info


CONFIG_PATH = pathlib.Path(__file__).parent / "config.ini"
RESULTS_DIR = pathlib.Path(__file__).parent / "results"


def get_api_key():
    """Return a usable API key, prompting and saving it on first run."""
    key = load_api_key(CONFIG_PATH)
    if key is not None:
        return key
    print("No API key saved yet.")
    print("Paste your OMIM API key (from the entitlement email) and press Enter.")
    key = input("API key: ").strip()
    while key == "":
        key = input("API key (cannot be blank): ").strip()
    save_api_key(CONFIG_PATH, key)
    print("Saved. You will not be asked again on this computer.")
    return key


def ask_guided_query():
    """Ask the guided questions and return an OMIM query string."""
    print()
    print("Enter search words. Separate several words with commas.")
    print("Leave a line blank to skip it.")
    any_of = split_terms(input("Find entries containing ANY of these words: "))
    must = split_terms(input("Words that MUST appear: "))
    exclude = split_terms(input("Words to EXCLUDE: "))
    return build_query(any_of, must, exclude)


def ask_max_results(total):
    """Ask how many results to fetch.

    Returns "cancel", a positive int cap, or None to mean fetch all. Invalid
    input is always re-prompted; it never silently falls through to "fetch
    all", since that would defeat the purpose of the quota guard.
    """
    estimated_requests = (total + 19) // 20
    print()
    print("OMIM found " + str(total) + " matching entries.")
    print("Fetching them all needs about " + str(estimated_requests) + " requests.")
    while True:
        answer = input("Fetch [a]ll, a [n]umber, or [c]ancel? ").strip().lower()
        if answer == "a":
            return None
        if answer == "c":
            return "cancel"
        if answer == "n":
            return _ask_result_count()
        print("Please type a, n, or c.")


def _ask_result_count():
    """Ask for a positive result count. Return the int, or "cancel"."""
    while True:
        number = input("How many? ").strip().lower()
        if number == "c":
            return "cancel"
        if number.isdigit() and int(number) >= 1:
            return int(number)
        print("Please enter a whole number of 1 or more, or c to cancel.")


def main():
    """Run one interactive OMIM search."""
    print("OMIM keyword search")
    print("===================")
    api_key = get_api_key()
    client = OmimClient(api_key)

    mode_answer = input("Guided search [g] or expert raw query [e]? ").strip().lower()
    if mode_answer == "e":
        query = input("Raw OMIM query: ").strip()
        mode = "expert"
    else:
        query = ask_guided_query()
        mode = "guided"

    print()
    print("Query: " + query)
    confirm = input("Run this search? [y/n] ").strip().lower()
    if confirm != "y":
        print("Cancelled.")
        return

    try:
        total = probe_count(client, query)
    except AuthFailed:
        print("The API key was rejected. Delete config.ini and try again.")
        return
    except QuotaExhausted:
        print("The API key's quota is exhausted. Try again later.")
        return
    except OmimError:
        print("Could not reach OMIM right now. Check your internet connection and try again.")
        return

    if total == 0:
        print("No results. Try fewer or broader words.")
        return

    decision = ask_max_results(total)
    if decision == "cancel":
        print("Cancelled.")
        return
    max_results = decision

    print("Fetching...")
    entries, raw_responses, complete = fetch_entries(client, query, max_results)

    entry_rows = []
    phenotype_rows = []
    rank = 1
    for entry in entries:
        entry_rows.append(entry_to_row(entry, rank))
        phenotype_rows.extend(entry_to_phenotype_rows(entry))
        rank = rank + 1

    search_info = build_search_info(
        query, mode, total, entry_rows, phenotype_rows, complete
    )
    folder = make_run_folder(RESULTS_DIR, query)
    write_run(folder, entry_rows, phenotype_rows, raw_responses, search_info)

    print()
    if not complete:
        print("NOTE: the run stopped early; results are partial (see Search info).")
    print("Saved " + str(len(entry_rows)) + " entries to:")
    print("  " + str(folder / "results.xlsx"))


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print()
        print("Stopped. Any results already fetched were not saved.")
        sys.exit(1)
