"""Count PubMed hits for each term in a list of condition names.

Reads a column of terms from an .xlsx or .csv file and, for each, asks PubMed
how many articles match - once as an exact quoted phrase and once as PubMed's
default search - so terms can be triaged before building a combined search.
"""

import configparser
import csv
import datetime
import pathlib
import sys
import time
import urllib.parse

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from omim_search import slugify, make_run_folder, write_csv


def read_terms(path):
    """Return the non-empty, trimmed first-column values from an xlsx or csv."""
    path = pathlib.Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        rows = _first_column_xlsx(path)
    else:
        rows = _first_column_csv(path)
    terms = []
    for value in rows:
        if value is None:
            continue
        text = str(value).strip()
        if text != "":
            terms.append(text)
    return terms


def _first_column_xlsx(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = []
    for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True):
        values.append(row[0])
    return values


def _first_column_csv(path):
    values = []
    with open(path, "r", encoding="utf-8-sig", newline="") as csv_file:
        reader = csv.reader(csv_file)
        for row in reader:
            if row:
                values.append(row[0])
            else:
                values.append(None)
    return values


PUBMED_WEB = "https://pubmed.ncbi.nlm.nih.gov/"


def exact_query(term):
    """Wrap the term in quotes so PubMed searches the exact phrase (no term mapping)."""
    return '"' + term.strip() + '"'


def natural_query(term):
    """Return the term unquoted, as a plain PubMed search would receive it."""
    return term.strip()


def pubmed_url(query):
    """Build the human-facing PubMed search URL for a query."""
    encoded = urllib.parse.urlencode({"term": query})
    return PUBMED_WEB + "?" + encoded


ESEARCH_URL = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi"


class PubMedError(Exception):
    """A PubMed request failed or returned an unusable response."""


def parse_count(payload):
    """Return the integer hit count from an ESearch JSON payload."""
    result = payload.get("esearchresult", {})
    if "ERROR" in result:
        raise PubMedError("PubMed returned an error: " + str(result["ERROR"]))
    count = result.get("count")
    if count is None:
        raise PubMedError("PubMed response had no count.")
    try:
        return int(count)
    except (TypeError, ValueError) as exc:
        raise PubMedError("PubMed count was not a number: " + str(count)) from exc


class PubMedClient:
    """A small, throttled, retrying client for PubMed ESearch count queries."""

    def __init__(self, email, api_key=None, tool="omim-pubmed-counter",
                 session=None, sleep=time.sleep):
        self.email = email
        self.api_key = api_key
        self.tool = tool
        if session is None:
            session = requests.Session()
        self.session = session
        self.sleep = sleep
        if api_key:
            self.pause_seconds = 0.11
        else:
            self.pause_seconds = 0.34

    def count(self, query):
        """Return how many PubMed records match the query string."""
        params = {
            "db": "pubmed",
            "term": query,
            "retmode": "json",
            "retmax": 0,
            "tool": self.tool,
            "email": self.email,
        }
        if self.api_key:
            params["api_key"] = self.api_key

        max_attempts = 3
        attempt = 1
        while True:
            self.sleep(self.pause_seconds)
            try:
                response = self.session.get(ESEARCH_URL, params=params, timeout=60)
            except requests.exceptions.RequestException as exc:
                if attempt >= max_attempts:
                    raise PubMedError("A network error occurred talking to PubMed.") from exc
                self.sleep(2 ** attempt)
                attempt = attempt + 1
                continue

            status = response.status_code
            if status == 200:
                return parse_count(response.json())
            if attempt >= max_attempts:
                raise PubMedError("PubMed returned HTTP " + str(status) + ".")
            self.sleep(2 ** attempt)
            attempt = attempt + 1
