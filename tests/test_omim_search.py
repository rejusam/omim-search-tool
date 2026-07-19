import csv
import json
import pathlib

import openpyxl

import omim_search

FIXTURE_PATH = pathlib.Path(__file__).parent / "fixtures" / "search_response.json"


def _load_entries():
    with open(FIXTURE_PATH, encoding="utf-8") as fixture_file:
        payload = json.load(fixture_file)
    wrapped = payload["omim"]["searchResponse"]["entryList"]
    return [item["entry"] for item in wrapped]


def test_load_api_key_missing_file_returns_none(tmp_path):
    missing = tmp_path / "config.ini"
    assert omim_search.load_api_key(missing) is None


def test_save_then_load_round_trips_key(tmp_path):
    config_path = tmp_path / "config.ini"
    omim_search.save_api_key(config_path, "ABC123")
    assert omim_search.load_api_key(config_path) == "ABC123"


def test_load_api_key_ignores_blank(tmp_path):
    config_path = tmp_path / "config.ini"
    omim_search.save_api_key(config_path, "   ")
    assert omim_search.load_api_key(config_path) is None


def test_quote_term_single_word_unquoted():
    assert omim_search.quote_term("neuropathy") == "neuropathy"


def test_quote_term_multi_word_quoted():
    assert omim_search.quote_term("muscular dystrophy") == '"muscular dystrophy"'


def test_build_query_any_of_wraps_in_required_group():
    query = omim_search.build_query(["neuropathy", "neuronopathy"], [], [])
    assert query == "+(neuropathy neuronopathy)"


def test_build_query_reference_case_has_no_stray_operator():
    # The whole reason guided mode exists: "neuropathy or neuronopathy"
    # must never become a query containing the term "or".
    query = omim_search.build_query(["neuropathy", "neuronopathy"], [], [])
    assert " or " not in query.lower()


def test_build_query_combines_all_three_groups():
    query = omim_search.build_query(["neuropathy"], ["hereditary"], ["diabetic"])
    assert query == "+(neuropathy) +hereditary -diabetic"


def test_build_query_must_include_multiword_is_quoted():
    query = omim_search.build_query([], ["muscular dystrophy"], [])
    assert query == '+"muscular dystrophy"'


def test_build_query_all_empty_raises():
    import pytest
    with pytest.raises(ValueError):
        omim_search.build_query([], [], [])


def test_page_starts_single_page():
    assert omim_search.page_starts(5, None) == [0]


def test_page_starts_exact_multiple():
    assert omim_search.page_starts(40, None) == [0, 20]


def test_page_starts_partial_last_page():
    assert omim_search.page_starts(45, None) == [0, 20, 40]


def test_page_starts_zero_results():
    assert omim_search.page_starts(0, None) == []


def test_page_starts_capped_by_max_results():
    assert omim_search.page_starts(1000, 30) == [0, 20]


def test_page_starts_max_larger_than_total():
    assert omim_search.page_starts(10, 500) == [0]


class FakeResponse:
    def __init__(self, status_code, payload=None):
        self.status_code = status_code
        self._payload = payload if payload is not None else {}

    def json(self):
        return self._payload


class FakeSession:
    def __init__(self, responses):
        # responses: a list of FakeResponse (or Exception instances to raise),
        # returned/raised in order
        self._responses = list(responses)
        self.calls = []

    def get(self, url, params=None, headers=None, timeout=None):
        self.calls.append({"url": url, "params": params, "headers": headers})
        item = self._responses.pop(0)
        if isinstance(item, Exception):
            raise item
        return item


def _client(responses):
    session = FakeSession(responses)
    client = omim_search.OmimClient(
        "TESTKEY", session=session, pause_seconds=0, sleep=lambda seconds: None
    )
    return client, session


def test_request_sends_api_key_header():
    client, session = _client([FakeResponse(200, {"omim": {"ok": True}})])
    client.request("status", {})
    assert session.calls[0]["headers"]["ApiKey"] == "TESTKEY"


def test_request_returns_omim_payload():
    client, _ = _client([FakeResponse(200, {"omim": {"value": 42}})])
    result = client.request("status", {})
    assert result == {"value": 42}


def test_request_401_raises_authfailed():
    import pytest
    client, _ = _client([FakeResponse(401)])
    with pytest.raises(omim_search.AuthFailed):
        client.request("status", {})


def test_request_429_raises_quotaexhausted_without_retry():
    import pytest
    client, session = _client([FakeResponse(429)])
    with pytest.raises(omim_search.QuotaExhausted):
        client.request("status", {})
    assert len(session.calls) == 1  # not retried


def test_request_500_retries_then_raises():
    import pytest
    client, session = _client([FakeResponse(500), FakeResponse(500), FakeResponse(500)])
    with pytest.raises(omim_search.OmimError):
        client.request("status", {})
    assert len(session.calls) == 3  # 3 attempts


def test_request_500_then_200_succeeds():
    client, session = _client([FakeResponse(500), FakeResponse(200, {"omim": {"ok": 1}})])
    result = client.request("status", {})
    assert result == {"ok": 1}
    assert len(session.calls) == 2


def test_request_connection_error_then_200_succeeds():
    import requests

    client, session = _client([
        requests.exceptions.ConnectionError("boom"),
        FakeResponse(200, {"omim": {"ok": 1}}),
    ])
    result = client.request("status", {})
    assert result == {"ok": 1}
    assert len(session.calls) == 2


def test_request_connection_error_always_raises_omimerror_after_three_attempts():
    import pytest
    import requests

    client, session = _client([
        requests.exceptions.ConnectionError("boom"),
        requests.exceptions.Timeout("timed out"),
        requests.exceptions.ConnectionError("boom again"),
    ])
    with pytest.raises(omim_search.OmimError):
        client.request("status", {})
    assert len(session.calls) == 3


def test_request_404_returns_empty_payload_without_retry():
    client, session = _client([FakeResponse(404)])
    result = client.request("status", {})
    assert result == {}
    assert len(session.calls) == 1


def test_search_builds_expected_path_and_params():
    client, session = _client([FakeResponse(200, {"omim": {"searchResponse": {}}})])
    client.search("+(neuropathy)", start=0, limit=20, include="geneMap")
    call = session.calls[0]
    assert call["url"].endswith("/entry/search")
    assert call["params"]["search"] == "+(neuropathy)"
    assert call["params"]["start"] == 0
    assert call["params"]["limit"] == 20
    assert call["params"]["include"] == "geneMap"
    assert call["params"]["format"] == "json"


# The fixture mirrors the real OMIM response shape: a phenotype entry (#)
# carries phenotypeMapList directly with gene fields embedded in each map;
# a gene entry (*) carries a geneMap with geneName plus a nested
# phenotypeMapList; a third entry has neither.


def test_entry_to_row_phenotype_entry_core_fields():
    entries = _load_entries()
    row = omim_search.entry_to_row(entries[0], rank=1)
    assert row["rank"] == 1
    assert row["mim_number"] == "100001"
    assert row["prefix"] == "#"
    assert row["status"] == "live"
    assert "EXAMPLE PHENOTYPE ENTRY" in row["preferred_title"]
    # Phenotype entries embed the gene fields in the phenotype map.
    assert row["gene_symbols"] == "GENEA, GENEA1, SYNA"
    assert row["approved_gene_symbol"] == "GENEA"
    assert row["cyto_location"] == "17p12"
    assert row["chromosome"] == "17"
    # geneName only exists on a geneMap, so a phenotype entry has none.
    assert row["gene_name"] == ""
    assert row["omim_url"] == "https://omim.org/entry/100001"


def test_entry_to_row_phenotype_summary():
    entries = _load_entries()
    row = omim_search.entry_to_row(entries[0], rank=1)
    assert row["phenotype_count"] == 2
    assert "Example phenotype, type 1A" in row["phenotypes"]
    assert " | " in row["phenotypes"]
    assert "Autosomal dominant" in row["inheritance"]
    assert "Autosomal recessive" in row["inheritance"]


def test_entry_to_row_gene_entry_reads_gene_map():
    entries = _load_entries()
    row = omim_search.entry_to_row(entries[1], rank=2)
    assert row["mim_number"] == "100002"
    assert row["prefix"] == "*"
    assert row["gene_symbols"] == "GENEB, GENEB1"
    assert row["approved_gene_symbol"] == "GENEB"
    assert row["gene_name"] == "Example gene B protein"
    assert row["cyto_location"] == "11q13"
    assert row["chromosome"] == "11"
    # Gene entries carry their phenotypes nested inside the geneMap.
    assert row["phenotype_count"] == 1
    assert row["phenotypes"] == "Example phenotype from gene B"


def test_entry_to_row_no_map_entry_has_blank_gene_fields():
    entries = _load_entries()
    row = omim_search.entry_to_row(entries[2], rank=3)
    assert row["mim_number"] == "100003"
    assert row["gene_symbols"] == ""
    assert row["approved_gene_symbol"] == ""
    assert row["gene_name"] == ""
    assert row["phenotype_count"] == 0
    assert row["phenotypes"] == ""


def test_entry_to_phenotype_rows_phenotype_entry():
    entries = _load_entries()
    rows = omim_search.entry_to_phenotype_rows(entries[0])
    assert len(rows) == 2
    assert rows[0]["phenotype"] == "Example phenotype, type 1A"
    assert rows[0]["phenotype_inheritance"] == "Autosomal dominant"
    assert rows[0]["phenotype_mim_number"] == "100001"
    assert rows[0]["phenotype_url"] == "https://omim.org/entry/100001"
    # Gene fields come from the phenotype map itself here.
    assert rows[0]["gene_symbols"] == "GENEA, GENEA1, SYNA"
    assert rows[0]["approved_gene_symbol"] == "GENEA"
    assert rows[1]["phenotype_inheritance"] == "Autosomal recessive"


def test_entry_to_phenotype_rows_gene_entry_uses_gene_map_fallback():
    entries = _load_entries()
    rows = omim_search.entry_to_phenotype_rows(entries[1])
    assert len(rows) == 1
    assert rows[0]["phenotype"] == "Example phenotype from gene B"
    assert rows[0]["phenotype_mim_number"] == "300003"
    assert rows[0]["phenotype_url"] == "https://omim.org/entry/300003"
    # The nested phenotype map has no gene fields, so they fall back
    # to the entry's geneMap.
    assert rows[0]["gene_symbols"] == "GENEB, GENEB1"
    assert rows[0]["approved_gene_symbol"] == "GENEB"
    assert rows[0]["cyto_location"] == "11q13"


def test_entry_to_phenotype_rows_no_map_entry_empty():
    entries = _load_entries()
    rows = omim_search.entry_to_phenotype_rows(entries[2])
    assert rows == []


def test_entry_to_row_missing_mim_number_has_empty_url():
    entry = {"titles": {"preferredTitle": "NO NUMBER ENTRY"}}
    row = omim_search.entry_to_row(entry, rank=1)
    assert row["mim_number"] == ""
    assert row["omim_url"] == ""


def test_entry_to_row_tolerates_null_field_values():
    # OMIM sometimes returns a field present but set to JSON null. This must
    # not crash the ' | ' join or leave None in a cell.
    entry = {
        "mimNumber": "100200",
        "prefix": "#",
        "status": "live",
        "titles": {"preferredTitle": "NULL FIELD ENTRY"},
        "phenotypeMapList": [
            {"phenotypeMap": {
                "phenotype": "Condition with unknown inheritance",
                "phenotypeInheritance": None,
                "phenotypeMimNumber": "100200",
            }},
        ],
    }
    row = omim_search.entry_to_row(entry, rank=1)
    assert row["inheritance"] == ""
    assert row["phenotypes"] == "Condition with unknown inheritance"
    rows = omim_search.entry_to_phenotype_rows(entry)
    assert rows[0]["phenotype_inheritance"] == ""


def test_entry_to_phenotype_rows_missing_phenotype_mim_has_empty_url():
    entry = {
        "mimNumber": "100100",
        "titles": {"preferredTitle": "TEST ENTRY"},
        "phenotypeMapList": [{
            "phenotypeMap": {"phenotype": "Some condition"},
        }],
    }
    rows = omim_search.entry_to_phenotype_rows(entry)
    assert len(rows) == 1
    assert rows[0]["phenotype_mim_number"] == ""
    assert rows[0]["phenotype_url"] == ""
    assert rows[0]["entry_url"] == "https://omim.org/entry/100100"


def _search_payload(total, entries):
    return {
        "omim": {
            "searchResponse": {
                "totalResults": total,
                "entryList": [{"entry": entry} for entry in entries],
            }
        }
    }


def test_probe_count_reads_total_results():
    client, _ = _client([FakeResponse(200, _search_payload(137, []))])
    assert omim_search.probe_count(client, "+(neuropathy)") == 137


def test_fetch_entries_single_page_complete():
    entries = [{"mimNumber": 1}, {"mimNumber": 2}]
    client, _ = _client([FakeResponse(200, _search_payload(2, entries))])
    got, raw, complete = omim_search.fetch_entries(client, "+(x)", max_results=None)
    assert [e["mimNumber"] for e in got] == [1, 2]
    assert complete is True
    assert len(raw) == 1


def test_fetch_entries_stops_on_quota_and_reports_incomplete():
    page_one = _search_payload(60, [{"mimNumber": 1}])
    client, _ = _client([
        FakeResponse(200, page_one),   # probe reuse not assumed; first fetch page
        FakeResponse(429),             # second page hits quota
    ])
    got, raw, complete = omim_search.fetch_entries(client, "+(x)", max_results=None)
    assert complete is False
    assert len(got) == 1  # partial results kept


def test_slugify_collapses_and_lowercases():
    assert omim_search.slugify("+(Neuropathy Neuronopathy)") == "neuropathy_neuronopathy"


def test_slugify_truncates_long_input():
    slug = omim_search.slugify("a" * 100)
    assert len(slug) <= 40


def test_make_run_folder_creates_unique_directory(tmp_path):
    first = omim_search.make_run_folder(tmp_path, "+(x)")
    second = omim_search.make_run_folder(tmp_path, "+(x)")
    assert first.exists()
    assert second.exists()
    assert first != second


def test_write_csv_writes_header_and_rows(tmp_path):
    path = tmp_path / "entries.csv"
    rows = [{"a": "1", "b": "SEPT9"}, {"a": "2", "b": "MARCH1"}]
    omim_search.write_csv(path, rows, ["a", "b"])
    with open(path, encoding="utf-8-sig", newline="") as csv_file:
        read_rows = list(csv.DictReader(csv_file))
    assert read_rows[0]["b"] == "SEPT9"
    assert read_rows[1]["b"] == "MARCH1"


def test_write_json_round_trips(tmp_path):
    path = tmp_path / "raw.json"
    omim_search.write_json(path, [{"page": 1}], {"query": "+(x)"})
    with open(path, encoding="utf-8") as json_file:
        loaded = json.load(json_file)
    assert loaded["metadata"]["query"] == "+(x)"
    assert loaded["responses"] == [{"page": 1}]


ENTRY_FIELDS = omim_search.ENTRY_FIELDS
PHENOTYPE_FIELDS = omim_search.PHENOTYPE_FIELDS


def _sample_rows():
    entry_rows = [{
        "rank": 1, "mim_number": 613008, "prefix": "#",
        "preferred_title": "EXAMPLE", "status": "live",
        "gene_symbols": "SEPT9", "approved_gene_symbol": "SEPT9",
        "gene_name": "Septin 9", "cyto_location": "17q25.3",
        "chromosome": "17", "phenotype_count": 1,
        "phenotypes": "Example phenotype", "inheritance": "Autosomal dominant",
        "omim_url": "https://omim.org/entry/613008",
    }]
    phenotype_rows = [{
        "mim_number": 613008, "preferred_title": "EXAMPLE",
        "gene_symbols": "SEPT9", "approved_gene_symbol": "SEPT9",
        "cyto_location": "17q25.3", "phenotype": "Example phenotype",
        "phenotype_mim_number": 613008, "phenotype_mapping_key": 3,
        "phenotype_inheritance": "Autosomal dominant",
        "phenotypic_series_number": "", "entry_url": "https://omim.org/entry/613008",
        "phenotype_url": "https://omim.org/entry/613008",
    }]
    return entry_rows, phenotype_rows


def test_write_xlsx_has_three_named_sheets(tmp_path):
    path = tmp_path / "results.xlsx"
    entry_rows, phenotype_rows = _sample_rows()
    omim_search.write_xlsx(path, entry_rows, phenotype_rows, {"Query": "+(x)"})
    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["Entries", "Phenotypes", "Search info"]


def test_write_xlsx_preserves_gene_symbol_as_text(tmp_path):
    path = tmp_path / "results.xlsx"
    entry_rows, phenotype_rows = _sample_rows()
    omim_search.write_xlsx(path, entry_rows, phenotype_rows, {"Query": "+(x)"})
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Entries"]
    header = [cell.value for cell in sheet[1]]
    symbol_col = header.index("approved_gene_symbol") + 1
    value_cell = sheet.cell(row=2, column=symbol_col)
    assert value_cell.value == "SEPT9"
    assert value_cell.number_format == "@"


def test_write_xlsx_omim_url_is_hyperlinked(tmp_path):
    path = tmp_path / "results.xlsx"
    entry_rows, phenotype_rows = _sample_rows()
    omim_search.write_xlsx(path, entry_rows, phenotype_rows, {"Query": "+(x)"})
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Entries"]
    header = [cell.value for cell in sheet[1]]
    url_col = header.index("omim_url") + 1
    cell = sheet.cell(row=2, column=url_col)
    assert cell.hyperlink is not None
    assert cell.hyperlink.target == "https://omim.org/entry/613008"


def test_write_xlsx_blank_url_is_not_hyperlinked(tmp_path):
    path = tmp_path / "results.xlsx"
    entry_rows, phenotype_rows = _sample_rows()
    entry_rows[0]["omim_url"] = ""
    omim_search.write_xlsx(path, entry_rows, phenotype_rows, {"Query": "+(x)"})
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Entries"]
    header = [cell.value for cell in sheet[1]]
    url_col = header.index("omim_url") + 1
    cell = sheet.cell(row=2, column=url_col)
    assert cell.hyperlink is None


def test_write_run_creates_all_four_files(tmp_path):
    folder = tmp_path / "run"
    folder.mkdir()
    entry_rows, phenotype_rows = _sample_rows()
    omim_search.write_run(folder, entry_rows, phenotype_rows, [{"page": 1}], {"Query": "+(x)"})
    assert (folder / "results.xlsx").exists()
    assert (folder / "entries.csv").exists()
    assert (folder / "phenotypes.csv").exists()
    assert (folder / "raw.json").exists()


def test_split_terms_splits_on_commas_and_trims():
    assert omim_search.split_terms("neuropathy, neuronopathy") == ["neuropathy", "neuronopathy"]


def test_split_terms_empty_string_is_empty_list():
    assert omim_search.split_terms("   ") == []


def test_build_search_info_reports_both_row_counts():
    info = omim_search.build_search_info(
        query="+(x)", mode="guided", total=50,
        entry_rows=[{}, {}, {}], phenotype_rows=[{}, {}], complete=True,
    )
    assert info["Entries written (Sheet 1)"] == 3
    assert info["Phenotype rows written (Sheet 2)"] == 2
    assert info["Total results reported by OMIM"] == 50
    assert info["Run complete"] == "yes"
    assert "not for redistribution" in info["Licence"].lower()
